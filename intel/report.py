"""Report generator: executive Markdown + Excel workbook with detail sheets."""
from __future__ import annotations

import sqlite3
from datetime import datetime
from pathlib import Path

import analytics
from normalize import format_inr

REPORT_DIR = Path("reports")


# ---------------------------------------------------------------------------
# Markdown executive report
# ---------------------------------------------------------------------------
def md_table(rows: list[list], headers: list[str]) -> str:
    out = ["| " + " | ".join(headers) + " |",
           "|" + "---|" * len(headers)]
    for r in rows:
        out.append("| " + " | ".join(str(c) if c is not None else "-" for c in r) + " |")
    return "\n".join(out)


def render_markdown(conn: sqlite3.Connection) -> str:
    from config import portal_label

    ms = analytics.market_size(conn)
    wm = analytics.winner_map(conn)
    opp = analytics.opportunity(conn, limit=30)
    buyers = analytics.buyer_profile(conn)
    caps = analytics.capacity_mix(conn)
    rates = analytics.rate_bench(conn)
    states = analytics.state_split(conn)
    portals = analytics.portal_split(conn)
    months = ms["by_month"]

    L = []
    L.append("# Water Cooler Govt Tender Intelligence Report")
    L.append(f"\n**Generated:** {datetime.now():%d-%b-%Y %H:%M}  |  **Coverage:** GeM + CPPP + manual award log")
    L.append("")
    L.append("## 1. Executive Summary")
    if ms["total"] != ms["relevant"]:
        L.append(f"- **Bids/tenders captured:** {ms['total']} total search hits, of which **{ms['relevant']} are relevant water-cooler bids** (noise filtered)  | currently open: {ms['open']}")
    else:
        L.append(f"- **Bids/tenders captured:** {ms['relevant']}  (currently open: {ms['open']})")
    if ms["est_value_sum"]:
        L.append(f"- **Est. value of tenders with known value ({ms['est_value_rows']} rows):** {format_inr(ms['est_value_sum'])}")
    L.append(f"- **Award log entries:** {wm['awards']}  |  **Units awarded:** {wm['qty_sold'] or 0:,.0f}")
    L.append(f"- **Seller types winning:** {', '.join(r['st'] + ' (' + str(r['awards']) + ')' for r in wm['by_type']) or 'none yet'}")
    L.append("")

    L.append("## 2. Why this matters (your levers)")
    L.append("- **Win more:** open bids at the top of this report are your hunting list; bid docs give exact spec/EMD/pBG so you can price to L1.")
    L.append("- **Market sizing:** est-value column is the demand signal per state and buyer; bigger than one tender = strategy.")
    L.append("- **Rate benchmark:** the award log tells you the L1 price Fresco Casa must beat per capacity class.")
    L.append("")

    L.append("## 3. Open Bid Pipeline (closing soon first)")
    if opp:
        L.append(md_table(
            [[r["bid_number"], (r["title"] or "")[:60],
              (r["org_name"] or r["ministry"]) or "-",
              r["end_date"], r["quantity"] or "-",
              r["capacity_class"] if r["capacity_class"] else "-",
              format_inr(r["est_value"]) if r["est_value"] else "-",
              portal_label(r["portal_code"] or "")] for r in opp],
            ["Bid No", "Title", "Buyer", "Closes", "Qty", "Cap(L)", "Est", "Source"],
        ))
    else:
        L.append("_No open bids in DB yet - run `python run.py scrape`._")
    L.append("")

    L.append("## 4. Market Volume by Month (captured)")
    L.append(md_table([[m["ym"], m["c"]] for m in months], ["Month", "Tenders"]));
    L.append("")

    L.append("## 5. Tenders by Capacity Class")
    L.append(md_table([[("-" if c["capacity_class"] is None else str(c["capacity_class"])),
                        c["tenders"], ("%s" % c["qty"]) if c["qty"] else "-"] for c in caps],
                      ["Capacity(L)", "Tenders", "Qty"]))
    L.append("")

    L.append("## 6. Top Buyers (by est. value)")
    L.append(md_table([[b["buyer"], b["segment"], b["tenders"],
                        format_inr(b["est_value"]) if b["est_value"] else "-",
                        b["qty"] or "-"] for b in buyers],
                      ["Buyer", "Segment", "Tenders", "Est. value", "Qty"]))
    L.append("")

    L.append("## 7. State Split")
    L.append(md_table([[s["st"], s["tenders"], format_inr(s["ev"]) if s["ev"] else "-"] for s in states],
                      ["State", "Tenders", "Est. value"]))
    L.append("")

    L.append("## 8. Winner Map (award log)")
    L.append(md_table([[w["seller_name"], w["seller_type"], w["awards"],
                        format_inr(w["value"]) if w["value"] else "-"] for w in wm["by_seller"]],
                      ["Seller", "Type", "Awards", "Value (INR)"]))
    L.append("")

    L.append("## 9. Rate Benchmark (per capacity, from award log)")
    L.append(md_table([[r["capacity_class"], r["n"], format_inr(r["lo"]), format_inr(r["avg"]),
                        format_inr(r["hi"]), r["flag"]] for r in rates],
                      ["Cap(L)", "n", "Low", "Avg", "High", "Band check"]))
    L.append("")

    L.append("## 10. Coverage by Source / Portal")
    L.append(md_table(
        [[portal_label(p["portal"]), p["portal"], p["n"], p["rel"] or 0,
          format_inr(p["ev"]) if p["ev"] else "-"] for p in portals],
        ["Portal", "Code", "Rows", "Water-cooler relevant", "Est. value"]))
    L.append("")
    L.append("---")
    L.append(f"*Data sources: GeM bidplus API, CPPP + {sum(1 for p in portals if p['portal'] not in ('gem','cppp'))} regional/state portals, manual award log.*")
    return "\n".join(L)


# ---------------------------------------------------------------------------
# Excel workbook export
# ---------------------------------------------------------------------------
def render_excel(conn: sqlite3.Connection, path: Path) -> None:
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Market Snapshot"
    header_font = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F4E78")

    def write_table(ws, headers, rows):
        for ci, h in enumerate(headers, 1):
            c = ws.cell(row=1, column=ci, value=h)
            c.font = header_font
            c.fill = fill
        for ri, row in enumerate(rows, 2):
            for ci, v in enumerate(row, 1):
                ws.cell(row=ri, column=ci, value=v)
        for ci in range(1, len(headers) + 1):
            ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = 18

    ms = analytics.market_size(conn)
    wm = analytics.winner_map(conn)
    write_table(ws,
                ["Metric", "Value"],
                [["Total bids captured", ms["total"]], ["Open now", ms["open"]],
                 ["Est. value (known rows)", format_inr(ms["est_value_sum"])],
                 ["Awards logged", wm["awards"]], ["Units awarded", wm["qty_sold"]]])

    sheets = {
        "Open Bids": (["Bid No", "Title", "Buyer", "Closes", "Qty", "Cap(L)", "Est"],
                      [[r["bid_number"], r["title"], r["org_name"] or r["ministry"],
                        r["end_date"], r["quantity"], r["capacity_class"],
                        format_inr(r["est_value"]) if r["est_value"] else "-"]
                       for r in analytics.opportunity(conn, limit=2000)]),
        "Buyers": (["Buyer", "Segment", "Tenders", "Est. value", "Qty"],
                   [[b["buyer"], b["segment"], b["tenders"],
                     format_inr(b["est_value"]) if b["est_value"] else "-", b["qty"]]
                    for b in analytics.buyer_profile(conn, limit=100)]),
        "Capacity Mix": (["Cap(L)", "Tenders", "Qty"],
                         [[c["capacity_class"], c["tenders"], c["qty"]] for c in analytics.capacity_mix(conn)]),
        "Rates": (["Cap(L)", "n", "Low", "Avg", "High", "Band check"],
                  [[r["capacity_class"], r["n"], format_inr(r["lo"]), format_inr(r["avg"]),
                    format_inr(r["hi"]), r["flag"]] for r in analytics.rate_bench(conn)]),
        "Winners": (["Seller", "Type", "Awards", "Value"],
                    [[w["seller_name"], w["seller_type"], w["awards"],
                      format_inr(w["value"]) if w["value"] else "-"] for w in wm["by_seller"]]),
        "State Split": (["State", "Tenders", "Est. value"],
                        [[s["st"], s["tenders"], format_inr(s["ev"]) if s["ev"] else "-"]
                         for s in analytics.state_split(conn)]),
        "By Source": (["Portal", "Code", "Rows", "Water-cooler relevant", "Est. value"],
                      [[r["portal"], r["portal"], r["n"], r["rel"], r["ev"]]
                       for r in analytics.portal_split(conn)]),
        "Raw Bids": (["Bid No", "Title", "Category", "Org", "Ministry", "Status",
                      "Cap(L)", "Qty", "Est", "Start", "End", "URL"],
                     [[r["bid_number"], r["title"], r["category"], r["org_name"],
                       r["ministry"], r["status"], r["capacity_class"], r["quantity"],
                       r["est_value"], r["start_date"], r["end_date"], r["source_url"]]
                      for r in conn.execute("SELECT * FROM bids ORDER BY end_date DESC LIMIT 5000")]),
        "Award Log": (["ID", "Date", "Buyer", "State", "Seller", "Type", "Cap(L)",
                       "Qty", "Unit", "Total", "Tender", "Source"],
                      [[r["award_id"], r["award_date"], r["buyer_org"], r["buyer_state"],
                        r["seller_name"], r["seller_type"], r["capacity_class"],
                        r["qty"], r["unit_price"], r["total_value"], r["tender_id"], r["source"]]
                       for r in conn.execute("SELECT * FROM awards ORDER BY award_date DESC LIMIT 5000")]),
    }
    for name, (headers, rows) in sheets.items():
        ws2 = wb.create_sheet(name)
        write_table(ws2, headers, rows)
    wb.save(path)


def build(conn: sqlite3.Connection, out_dir: Path | None = None) -> dict:
    out_dir = out_dir or REPORT_DIR
    out_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    md_path = out_dir / f"water_cooler_report_{ts}.md"
    xlsx_path = out_dir / f"water_cooler_report_{ts}.xlsx"
    md_path.write_text(render_markdown(conn), encoding="utf-8")
    render_excel(conn, xlsx_path)
    return {"markdown": md_path, "excel": xlsx_path}